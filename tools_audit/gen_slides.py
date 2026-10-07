# -*- coding: utf-8 -*-
"""Regenerates the 12 defense slides (dark theme) with data consistent with the portal DB,
then packs them into reports/lab7/LR7_Presentation_Variant15.{pptx,pdf}.

All text is drawn through helpers that word-wrap / shrink-to-fit and assert that nothing
leaves its container (card or slide), so overflow regressions fail the build instead of
silently producing broken slides.
"""
import os
from PIL import Image, ImageDraw, ImageFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FD = '/usr/share/fonts/truetype/dejavu/'
_FONTS = {}


def F(sz, bold=False):
    key = (sz, bold)
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(FD + ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf'), sz)
    return _FONTS[key]


BG = '#111827'; PANEL = '#1f2937'; ACC = '#2563eb'; SKY = '#38bdf8'; TXT = '#f9fafb'; MUT = '#9ca3af'
GREEN = '#34d399'; PURP = '#a78bfa'; AMBER = '#fbbf24'; RED = '#f87171'; CHIP = '#374151'
W, H = 1920, 1080
MARGIN = 60                 # left/right safe margin
CONTENT_BOTTOM = H - 80     # nothing below this line except the footer
OUT = f'{ROOT}/diagrams/lab7'
os.makedirs(OUT, exist_ok=True)

# ----------------------------------------------------------------------------- helpers

class Overflow(Exception):
    pass


def tw(text, font):
    return ImageDraw.Draw(Image.new('RGB', (1, 1))).textlength(text, font=font)


def wrap(text, font, maxw):
    """Word-wrap `text` so that every line fits into `maxw` pixels."""
    words = text.split(' ')
    lines, cur = [], ''
    for w in words:
        cand = (cur + ' ' + w).strip()
        if tw(cand, font) <= maxw or not cur:
            cur = cand
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    for ln in lines:
        if tw(ln, font) > maxw:
            raise Overflow(f'word too long for {maxw}px: {ln!r}')
    return lines


def fit_font(text, maxw, size, bold=False, min_size=12):
    """Largest font size <= `size` at which `text` fits into `maxw`."""
    while size > min_size and tw(text, F(size, bold)) > maxw:
        size -= 1
    if tw(text, F(size, bold)) > maxw:
        raise Overflow(f'cannot fit {text!r} into {maxw}px')
    return F(size, bold)


def T(d, x, y, text, font, fill, maxw=None, maxy=CONTENT_BOTTOM, anchor=None):
    """Draw a single line and assert that it stays inside [x, x+maxw] x [.., maxy]."""
    if maxw is None:
        maxw = W - MARGIN - x
    width = tw(text, font)
    if width > maxw + 0.5:
        raise Overflow(f'text {text!r} ({width:.0f}px) wider than {maxw}px')
    if y + font.size * 1.25 > maxy:
        raise Overflow(f'text {text!r} below bottom limit {maxy}')
    d.text((x, y), text, font=font, fill=fill, anchor=anchor)
    return width


def para(d, x, y, text, font, fill, maxw, lh=None, maxy=CONTENT_BOTTOM):
    """Draw wrapped paragraph, return y after the last line."""
    lh = lh or int(font.size * 1.45)
    for ln in wrap(text, font, maxw):
        T(d, x, y, ln, font, fill, maxw, maxy)
        y += lh
    return y


def center(d, y, text, font, fill, cx=W // 2, maxw=W - 2 * MARGIN):
    width = tw(text, font)
    if width > maxw:
        raise Overflow(f'centered text {text!r} wider than {maxw}px')
    d.text((cx - width / 2, y), text, font=font, fill=fill)


def new_slide(num):
    im = Image.new('RGB', (W, H), BG); d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 8], fill=ACC); d.rectangle([0, H - 8, W, H], fill=ACC)
    T(d, MARGIN, H - 52, 'Корпоративный портал университета — Модуль Центра АНОК | 2026', F(16), MUT, maxy=H)
    lab = f'Слайд {num} из 12'
    T(d, W - MARGIN - tw(lab, F(16, True)), H - 52, lab, F(16, True), SKY, maxy=H)
    return im, d


def head(d, title, sub):
    badge_w, badge_x = 460, W - MARGIN - 460
    T(d, MARGIN, 50, title, fit_font(title, badge_x - 30 - MARGIN, 34, True), TXT, badge_x - 30 - MARGIN)
    T(d, MARGIN, 100, sub, fit_font(sub, badge_x - 30 - MARGIN, 18), MUT, badge_x - 30 - MARGIN)
    d.rounded_rectangle([badge_x, 50, W - MARGIN, 110], 10, fill=PANEL)
    T(d, badge_x + 20, 68, 'ПРКП · Вариант 15 · Центр АНОК', F(16), MUT, badge_w - 40)
    d.rectangle([0, 130, W, 136], fill=ACC)


def card(d, x, y, w, h, border=ACC):
    if x < MARGIN - 1 or x + w > W - MARGIN + 1 or y + h > CONTENT_BOTTOM + 1:
        raise Overflow(f'card {x},{y},{w},{h} leaves the safe area')
    d.rounded_rectangle([x, y, x + w, y + h], 12, fill=PANEL)
    d.rectangle([x, y, x + w, y + 6], fill=border)
    return x + w, y + h  # right / bottom limits for content


def bullets(d, x, y, items, maxw, sz=20, col=TXT, step=None, maxy=CONTENT_BOTTOM):
    """Bulleted list with word-wrapping of long items; returns y after the last line."""
    font = F(sz); step = step or int(sz * 2.2); lh = int(sz * 1.4)
    for t in items:
        lines = wrap(t, font, maxw - 26)
        d.ellipse([x, y + sz * 0.45, x + 10, y + sz * 0.45 + 10], fill=SKY)
        yy = y
        for ln in lines:
            T(d, x + 26, yy, ln, font, col, maxw - 26, maxy)
            yy += lh
        y = yy + (step - lh)
    return y


def bullets_height(items, maxw, sz=20, step=None):
    font = F(sz); step = step or int(sz * 2.2); lh = int(sz * 1.4); h = 0
    for t in items:
        h += lh * len(wrap(t, font, maxw - 26)) + (step - lh)
    return h


def chips(d, x, y, items, maxw, font, pad=20, gap=16, row=56, fill=CHIP, col=TXT, maxy=CONTENT_BOTTOM):
    """Flow layout of pill-shaped chips, wrapping to the next row when needed."""
    x0 = x; cx = x
    for a in items:
        w = tw(a, font) + 2 * pad
        if w > maxw:
            raise Overflow(f'chip {a!r} wider than {maxw}px')
        if cx + w > x0 + maxw:
            cx = x0; y += row
        if y + 40 > maxy:
            raise Overflow(f'chip {a!r} below bottom limit {maxy}')
        d.rounded_rectangle([cx, y, cx + w, y + 40], 20, fill=fill)
        T(d, cx + pad, y + 10, a, font, col, w - 2 * pad, maxy)
        cx += w + gap
    return y + 40


def chips_height(items, maxw, font, pad=20, gap=16, row=56):
    cx = 0; rows = 1
    for a in items:
        w = tw(a, font) + 2 * pad
        if cx + w > maxw:
            cx = 0; rows += 1
        cx += w + gap
    return 40 + (rows - 1) * row


# ----------------------------------------------------------------------------- S1 title
im, d = new_slide(1)
badge = 'ИТОГОВАЯ ЗАЩИТА ПРОЕКТА КОРПОРАТИВНОГО ПОРТАЛА'
bf = F(17, True); bw = tw(badge, bf) + 80
d.rounded_rectangle([W / 2 - bw / 2, 150, W / 2 + bw / 2, 205], 26, fill=ACC)
center(d, 167, badge, bf, '#fff')
center(d, 260, 'Проектирование и разработка корпоративного веб-портала', F(40, True), TXT)
center(d, 330, '«Модуль Центра аккредитации и независимой оценки', F(36, True), SKY)
center(d, 395, 'качества образования (АНОК)»', F(36, True), SKY)
cw, gap = 420, 40
xs = [MARGIN + i * (cw + gap) for i in range(4)]
for x, (t, b, c) in zip(xs, [('Нормативная база', 'ФГОС 3++, матрицы ЗЕТ, лимиты часов и компетенций', ACC),
                             ('Сквозной воркфлоу', 'Согласование УП, служебные записки, протоколы УС', PURP),
                             ('Криптография ЭЦП', 'Хеширование SHA-256, проверка УЦ, штампы', GREEN),
                             ('Интерактивный UX', 'Chart.js, FullCalendar, Diff-трекер версий, SheetJS', SKY)]):
    card(d, x, 520, cw, 210, c)
    T(d, x + 22, 550, t, F(22, True), TXT, cw - 44)
    for i, ln in enumerate(b.split(', ')):
        T(d, x + 22, 600 + i * 30, ln, F(16), MUT, cw - 44)
T(d, MARGIN, 880, 'Дисциплина: Проектирование и разработка корпоративных порталов (ПРКП)', F(18), MUT, 1100)
T(d, MARGIN, 915, 'Специализация: Вариант № 15 (Университет — Центр АНОК)', F(18, True), SKY, 1100)
T(d, 1300, 880, 'Выполнил: студент группы П-41', F(22), MUT)
T(d, 1300, 915, 'Год разработки: 2026 · г. Москва', F(18), MUT)
im.save(f'{OUT}/slide_01_title.png')

# ----------------------------------------------------------------------------- S2 context & goals
im, d = new_slide(2); head(d, 'Контекст и цели проекта', 'Предметная область, акторы и нормативные требования')
left = ['Центр АНОК: экспертиза и валидация УП по ФГОС ВО 3++', 'Объект автоматизации — учебный план (240 з.е., 8 семестров)',
        'Маршрут утверждения: каф. → АНОК → Проректор → Ректор', 'Служебные записки на корректировку дисциплин («Было → Стало»)']
right = ['Алгоритмическая валидация нормативов (240 / 30 / ≤5 / ≥12 / ≥6)', 'Исключение ошибок несбалансированности з.е.',
         'Прозрачный жизненный цикл утверждения с ЭЦП (SHA-256)', 'Единое информационное пространство кафедр']
cw = (W - 2 * MARGIN - 30) // 2; inner = cw - 50
ch = 80 + max(bullets_height(left, inner, 19), bullets_height(right, inner, 19)) + 10
card(d, MARGIN, 170, cw, ch, ACC)
T(d, MARGIN + 25, 200, 'Предметная область', F(24, True), SKY, inner)
bullets(d, MARGIN + 25, 250, left, inner, 19)
x2 = MARGIN + cw + 30
card(d, x2, 170, cw, ch, GREEN)
T(d, x2 + 25, 200, 'Цели разработки', F(24, True), GREEN, inner)
bullets(d, x2 + 25, 250, right, inner, 19)
actors = ['Зав. выпускающей каф.', 'Зав. реализующей каф.', 'Эксперт АНОК', 'Начальник АНОК', 'Проректор по УР', 'Ректор',
          'Секретарь УС', 'Преподаватель / студент', 'Администратор']
ay = 170 + ch + 30; aw = W - 2 * MARGIN; af = F(17, True)
ah = 100 + chips_height(actors, aw - 50, af) + 30
card(d, MARGIN, ay, aw, ah, PURP)
T(d, MARGIN + 25, ay + 30, 'Основные акторы (9 ролей)', F(24, True), PURP, aw - 50)
chips(d, MARGIN + 25, ay + 100, actors, aw - 50, af)
ny = ay + ah + 30
norms = [('240 з.е.', 'общая трудоемкость ОП', ACC), ('30 з.е.', 'в каждом из 8 семестров', SKY), ('≤ 5', 'экзаменов в сессию', AMBER),
         ('≥ 12 з.е.', 'практическая подготовка (Б2)', GREEN), ('≥ 6 з.е.', 'блок ГИА (Б3)', PURP)]
n = len(norms); gap = 25; cw5 = (W - 2 * MARGIN - gap * (n - 1)) // n; x = MARGIN
T(d, MARGIN, ny, 'Нормативные требования ФГОС ВО 3++, проверяемые порталом', F(22, True), AMBER)
for v, lab, c in norms:
    card(d, x, ny + 44, cw5, 130, c)
    T(d, x + 26, ny + 70, v, F(34, True), c, cw5 - 52)
    T(d, x + 26, ny + 126, lab, F(16), MUT, cw5 - 52)
    x += cw5 + gap
im.save(f'{OUT}/slide_02_context_goals.png')

# ----------------------------------------------------------------------------- S3 architecture
im, d = new_slide(3); head(d, 'Архитектура и стек технологий', 'Многоуровневая архитектура без тяжеловесных фреймворков')
layers = [('Web UI', 'HTML5/CSS3, адаптив, JS-модули (Chart.js, FullCalendar, DataTables, SheetJS, jsPDF, CryptoJS, SweetAlert2/Toastify, Diff2Html)', ACC),
          ('HTTP Router', '15 разделов + REST API (/api/stats, /api/plans, /api/announcements)', SKY),
          ('Service Layer', 'CurriculumService · ValidationEngine · WorkflowManager · MemoManager · AuditLogger', PURP),
          ('Data Access', 'PDO MySQL / sqlite3, Prepared Statements; СУБД MySQL 8 / MariaDB / SQLite (13 таблиц, FK)', GREEN)]
lx, lw = 160, W - 320; inner = lw - 60; y = 180
for i, (name, desc, c) in enumerate(layers):
    lines = wrap(desc, F(21), inner)
    h = 34 + 40 + len(lines) * 32 + 24
    card(d, lx, y, lw, h, c)
    T(d, lx + 30, y + 28, f'{i + 1}. {name}', F(24, True), c, inner)
    yy = y + 70
    for ln in lines:
        T(d, lx + 30, yy, ln, F(21), TXT, inner); yy += 32
    if i < len(layers) - 1:
        d.polygon([(lx + lw // 2 - 14, y + h + 4), (lx + lw // 2 + 14, y + h + 4), (lx + lw // 2, y + h + 20)], fill=MUT)
    y += h + 26
metrics = [('< 50 мс', 'отклик страниц', GREEN), ('< 30 МБ', 'потребление RAM', SKY), ('python3 server.py', 'запуск одной командой', AMBER)]
n = len(metrics); gap = 30; mw = (lw - gap * (n - 1)) // n; x = lx; y += 6
for v, lab, c in metrics:
    card(d, x, y, mw, 110, c)
    T(d, x + 26, y + 26, v, F(30, True), c, mw - 52)
    T(d, x + 26, y + 74, lab, F(16), MUT, mw - 52)
    x += mw + gap
im.save(f'{OUT}/slide_03_architecture_stack.png')

# ----------------------------------------------------------------------------- S4 data model
im, d = new_slide(4); head(d, 'Информационная модель и структура базы данных', '13 нормализованных таблиц (3NF) с FOREIGN KEYS')
rows = [('curriculums', 'Реестр учебных планов', '1'), ('curriculum_disciplines', 'Семестровая матрица дисциплин', '51'),
        ('standards_fgos', 'Нормативы ФГОС 3++', '3'), ('educational_programs', 'Паспорта ОП', '1'), ('competencies', 'Матрица УК/ОПК/ПК', '13'),
        ('departments', 'Подразделения (выпускающие/обеспечивающие/АНОК)', '9'), ('document_signatures', 'Этапы визирования + SHA-256', '4'),
        ('service_notes', 'Служебные записки', '4'), ('service_note_items', 'Диффы «Было → Стало»', '4'), ('users', 'Учетные записи', '7'),
        ('roles', 'Роли доступа', '7'), ('announcements', 'Регламенты и объявления', '4'), ('audit_logs', 'Журнал аудита с IP', '7')]
cw = (W - 2 * MARGIN - 30) // 2; rh, step = 92, 104
for i, (nm, desc, cnt) in enumerate(rows):
    col = i % 2; xx = MARGIN + col * (cw + 30); yy = 170 + (i // 2) * step
    card(d, xx, yy, cw, rh, CHIP)
    T(d, xx + 24, yy + 16, nm, F(20, True), SKY, cw - 150)
    T(d, xx + 24, yy + 50, desc, F(16), MUT, cw - 150)
    bx = xx + cw - 110
    d.rounded_rectangle([bx, yy + 20, bx + 80, yy + 72], 8, fill=CHIP, outline=ACC)
    T(d, bx + 40 - tw(cnt, F(24, True)) / 2, yy + 30, cnt, F(24, True), TXT, 80)
total = sum(int(r[2]) for r in rows)
xx = MARGIN + cw + 30; yy = 170 + 6 * step
card(d, xx, yy, cw, rh, ACC)
T(d, xx + 24, yy + 16, f'Итого: {len(rows)} таблиц', F(20, True), TXT, cw - 150)
T(d, xx + 24, yy + 50, 'записей в seed.sql', F(16), MUT, cw - 150)
bx = xx + cw - 110
d.rounded_rectangle([bx, yy + 20, bx + 80, yy + 72], 8, fill=ACC)
T(d, bx + 40 - tw(str(total), F(24, True)) / 2, yy + 30, str(total), F(24, True), '#fff', 80)
fy = 170 + 7 * step - step + rh + 24
card(d, MARGIN, fy, W - 2 * MARGIN, 64, GREEN)
T(d, MARGIN + 30, fy + 20, 'Каскадная целостность: FOREIGN KEYS в schema.sql; PRAGMA foreign_key_check = 0 нарушений', F(19, True), GREEN, W - 2 * MARGIN - 60)
im.save(f'{OUT}/slide_04_data_model.png')

# ----------------------------------------------------------------------------- S5 structure
im, d = new_slide(5); head(d, 'Структура портала и навигация', '15 функциональных страниц + сквозная семестровая навигация')
groups = [('Учебные планы', ['dashboard', 'plans', 'plan (сем. 1–8)', 'competencies'], ACC),
          ('Качество / ФГОС', ['validate', 'standards', 'analytics (Chart.js)'], GREEN),
          ('Документооборот', ['signatures (ЭЦП)', 'memos (Diff)', 'calendar (FullCalendar)'], PURP),
          ('Справочники', ['programs', 'departments', 'announcements'], SKY),
          ('Служебные', ['audit', 'sitemap', 'REST API ×3'], AMBER)]
n = len(groups); gap = 25; cw = (W - 2 * MARGIN - gap * (n - 1)) // n
ch = 90 + 72 * max(len(g[1]) for g in groups) + 10
x = MARGIN
for g, items, c in groups:
    card(d, x, 180, cw, ch, c)
    T(d, x + 22, 210, g, F(22, True), c, cw - 44)
    yy = 270
    for it in items:
        d.rounded_rectangle([x + 22, yy, x + cw - 22, yy + 56], 8, fill=CHIP)
        label = it if it.startswith('REST') else '?page=' + it
        T(d, x + 38, yy + 17, label, fit_font(label, cw - 76, 16), TXT, cw - 76)
        yy += 72
    x += cw + gap
by = 180 + ch + 40
sem_items = [f'Семестр {i} · 30 з.е.' for i in range(1, 9)]
sem_font = F(16, True); sem_w = W - 2 * MARGIN - 60
nh = 160 + chips_height(sem_items, sem_w, sem_font, pad=16, gap=12) + 30
card(d, MARGIN, by, W - 2 * MARGIN, nh, SKY)
T(d, MARGIN + 30, by + 30, 'Навигация', F(22, True), SKY)
T(d, MARGIN + 30, by + 70, 'Сайдбар: 14 пунктов меню + страница просмотра плана · адаптив ≤ 768px · футер со статусом синхронизации', F(19), MUT, sem_w)
T(d, MARGIN + 30, by + 118, 'Сквозная семестровая навигация на странице плана (?page=plan&id=1&sem=N), баланс 30 з.е. / 1080 ч в каждом семестре:', F(17), TXT, sem_w)
chips(d, MARGIN + 30, by + 160, sem_items, sem_w, sem_font, pad=16, gap=12, maxy=by + nh)
im.save(f'{OUT}/slide_05_portal_structure.png')

# ----------------------------------------------------------------------------- S6 fgos
im, d = new_slide(6); head(d, 'Автоматизированная экспертиза ФГОС ВО 3++', 'Все показатели вычисляются из БД (функция validate_curriculum)')
hdr = ['Критерий', 'Норма', 'Факт', 'Статус']
rows = [('Общая трудоемкость', '240 з.е.', '240 з.е.', 'СООТВЕТСТВУЕТ'), ('Баланс семестров 1–8', '30 з.е. / 1080 ч', '30 / 1080 в каждом', 'СООТВЕТСТВУЕТ'),
        ('Экзаменов в сессию', '≤ 5', '4', 'СООТВЕТСТВУЕТ'), ('Практическая подготовка (Б2)', '≥ 12 з.е.', '34 з.е.', 'СООТВЕТСТВУЕТ'),
        ('Блок ГИА (Б3)', '≥ 6 з.е.', '9 з.е.', 'СООТВЕТСТВУЕТ')]
tx, tw_ = 80, W - 160
cols = [tx + 20, tx + 640, tx + 1020, tx + 1440]
widths = [600, 360, 400, tw_ - 1440 - 40]
y = 200
for xx, htxt, wd in zip(cols, hdr, widths):
    T(d, xx, y, htxt, F(20, True), SKY, wd)
y += 50
for r in rows:
    card(d, tx, y, tw_, 90, CHIP)
    for xx, val, wd in zip(cols, r, widths):
        ok = val == 'СООТВЕТСТВУЕТ'
        T(d, xx, y + 32, val, F(19, ok), GREEN if ok else TXT, wd)
    y += 110
card(d, tx, y + 10, tw_, 70, AMBER)
T(d, tx + 20, y + 33, 'Блоки УП: Б1.О 117 · Б1.В 71 · Б2 34 · Б3 9 · ФТД 9 з.е. = 240', F(20, True), AMBER, tw_ - 40)
im.save(f'{OUT}/slide_06_fgos_validation.png')

# ----------------------------------------------------------------------------- S7 memos
im, d = new_slide(7); head(d, 'Управление изменениями УП', 'Служебные записки с версионными изменениями «Было → Стало»')
rows = [('СЗ-2026/048', 'Матанализ: 3 → 5 з.е., зачет → экзамен', 'ПРИМЕНЕНА', GREEN),
        ('СЗ-2026/052', 'Алгоритмизация: 36 ч практики → 36 ч лабораторных', 'ПРИМЕНЕНА', GREEN),
        ('СЗ-2026/059', 'Мобильная разработка → Каф. ИСТ', 'НА ЭКСПЕРТИЗЕ', PURP),
        ('СЗ-2026/061', 'График преддипломной практики: сдвиг на 2 недели', 'ВИЗА ВЫПУСКАЮЩЕЙ', SKY)]
rx, rw = 80, W - 160; y = 180
for nm, desc, st, c in rows:
    card(d, rx, y, rw, 108, c)
    T(d, rx + 30, y + 40, nm, F(22, True), SKY, 240)
    T(d, rx + 300, y + 42, desc, F(20), TXT, rw - 300 - 420)
    sw = tw(st, F(19, True))
    d.rounded_rectangle([rx + rw - 30 - sw - 40, y + 32, rx + rw - 30, y + 76], 10, fill=CHIP)
    T(d, rx + rw - 30 - sw - 20, y + 42, st, F(19, True), c, sw + 2)
    y += 128
y += 14
card(d, rx, y, rw, 190, RED)
T(d, rx + 30, y + 30, 'Diff2Html (side-by-side): СЗ-2026/048', F(20, True), RED, rw - 60)
d.rounded_rectangle([rx + 30, y + 72, rx + rw - 30, y + 110], 6, fill='#3b1f24')
T(d, rx + 44, y + 81, '- Семестр 1: Матанализ — 3 з.е. (108 ч.) · Зачет · Каф. ВМ', F(18), RED, rw - 90)
d.rounded_rectangle([rx + 30, y + 122, rx + rw - 30, y + 160], 6, fill='#1b3a30')
T(d, rx + 44, y + 131, '+ Семестр 1: Матанализ — 5 з.е. (180 ч.) [+2 з.е.] · Экзамен · Каф. ВМ', F(18), GREEN, rw - 90)
im.save(f'{OUT}/slide_07_visual_diff.png')

# ----------------------------------------------------------------------------- S8 signatures
im, d = new_slide(8); head(d, 'Цифровое визирование и ЭЦП', '4 этапа согласования с криптографическими штампами SHA-256')
steps = [('1', 'Руководитель выпускающего подразделения', 'Иванов И.И. (Зав. каф. ПИ)', '2026-06-15', 'da15dc1a…ede6'),
         ('2', 'Начальник Центра АНОК', 'Кузнецов В.П. (Нач. АНОК)', '2026-06-18', '2e4aab7…6dc9'),
         ('3', 'Проректор по учебной работе', 'Смирнов А.Н. (Проректор по УР)', '2026-06-22', '67c1ef7…1f2d'),
         ('4', 'Ректор университета', 'Михайлов С.В. (Ректор) · протокол УС № 8/26 от 25.06.2026', '2026-06-25', '36f5e44…184')]
sx, sw_ = 140, W - 280; y = 180
for n, role, who, dt, hh in steps:
    card(d, sx, y, sw_, 136, GREEN)
    d.ellipse([sx + 30, y + 38, sx + 90, y + 98], fill=ACC)
    T(d, sx + 60 - tw(n, F(26, True)) / 2, y + 52, n, F(26, True), '#fff', 60)
    T(d, sx + 120, y + 22, role, F(22, True), TXT, sw_ - 150)
    T(d, sx + 120, y + 60, f'{who} · {dt}', F(17), MUT, sw_ - 150)
    T(d, sx + 120, y + 94, f'SHA-256: {hh} · [MATCH] (пересчет совпадает с хранилищем)', F(15), GREEN, sw_ - 150)
    y += 156
para(d, sx, y + 12, 'Верификация на странице ?page=signatures: CryptoJS пересчитывает payload и сверяет с БД (4 из 4 действительны)', F(18), MUT, sw_)
im.save(f'{OUT}/slide_08_crypto_signature.png')

# ----------------------------------------------------------------------------- S9 analytics
im, d = new_slide(9); head(d, 'Аналитика и сводная панель', 'Chart.js: структура блоков, компетенции, часы по семестрам')
vals = [('117', 'Б1.О обязательная', ACC), ('71', 'Б1.В вариативная', SKY), ('34', 'Б2 практики', GREEN), ('9', 'Б3 ГИА', AMBER), ('9', 'ФТД', PURP)]
n = len(vals); gap = 25; cw = (W - 2 * MARGIN - gap * (n - 1)) // n; x = MARGIN
for v, lab, c in vals:
    card(d, x, 180, cw, 180, c)
    T(d, x + 30, 212, v, F(46, True), c, cw - 60)
    T(d, x + 30, 290, lab + ' з.е.', F(18), MUT, cw - 60)
    x += cw + gap
left = ['Учебные планы: 1 · Дисциплины: 51', 'Стандарты ФГОС: 3 · Служебные записки: 4', 'Реестр УП с переходом к семестровой сетке', 'Лента закрепленных регламентов АНОК']
right = ['Doughnut: структура блоков (117/71/34/9/9)', 'Radar: покрытие компетенций УК 5 · ОПК 4 · ПК 4', 'Stacked Bar: 1080 ч в каждом семестре',
         'Данные диаграмм вычисляются на сервере из БД, а не хранятся константами']
cw = (W - 2 * MARGIN - 30) // 2; inner = cw - 60
ch = 100 + max(bullets_height(left, inner, 20, 48), bullets_height(right, inner, 20, 48)) + 10
card(d, MARGIN, 400, cw, ch, ACC)
T(d, MARGIN + 30, 432, 'Сводная панель', F(24, True), SKY, inner)
bullets(d, MARGIN + 30, 486, left, inner, 20, step=48)
x2 = MARGIN + cw + 30
card(d, x2, 400, cw, ch, PURP)
T(d, x2 + 30, 432, 'Диаграммы Chart.js', F(24, True), PURP, inner)
bullets(d, x2 + 30, 486, right, inner, 20, step=48)
ey = 400 + ch + 30
card(d, MARGIN, ey, W - 2 * MARGIN, 120, GREEN)
T(d, MARGIN + 30, ey + 30, 'Экспорт страницы плана', F(24, True), GREEN)
T(d, MARGIN + 30, ey + 72, 'SheetJS → XLSX семестровой сетки · jsPDF → PDF-выписка из учебного плана · window.print() → печать бланка УП', F(19), MUT, W - 2 * MARGIN - 60)
im.save(f'{OUT}/slide_09_analytics_export.png')

# ----------------------------------------------------------------------------- S10 ux
im, d = new_slide(10); head(d, 'Пользовательский интерфейс', 'Адаптивный дизайн, единая палитра, готовые модули')
tok = [('#0f172a', 'header'), ('#1e293b', 'sidebar'), ('#f8fafc', 'page'), ('#2563eb', 'primary'), ('#15803d', 'ok'), ('#b91c1c', 'err'), ('#7e22ce', 'anok'), ('#b45309', 'warn')]
n = len(tok); gap = 24; sw_ = (W - 2 * MARGIN - gap * (n - 1)) // n; x = MARGIN
for c, nm in tok:
    d.rounded_rectangle([x, 180, x + sw_, 260], 10, fill=c, outline='#475569')
    T(d, x + 12, 272, nm, F(15), MUT, sw_ - 12)
    T(d, x + 12, 294, c, F(13), '#6b7280', sw_ - 12)
    x += sw_ + gap
left = ['Header 60px sticky + Sidebar 240px + Main 1200px + Footer', 'Типографика: 24/18/14/12px, системный стек шрифтов',
        'Адаптив ≤ 768px: сайдбар 70px, overflow-x таблиц', '14 пунктов меню, подсветка активного раздела']
right = ['DataTables: живой поиск и пагинация плана', 'SheetJS / jsPDF: экспорт XLSX и PDF', 'FullCalendar: график процедур АНОК',
         'Diff2Html: сравнение версий УП', 'SweetAlert2 / Toastify: диалоги и уведомления']
cw = (W - 2 * MARGIN - 30) // 2; inner = cw - 60
ch = 90 + max(bullets_height(left, inner, 19), bullets_height(right, inner, 19)) + 10
card(d, MARGIN, 350, cw, ch, SKY)
T(d, MARGIN + 30, 380, 'Шаблон дизайна', F(22, True), SKY, inner)
bullets(d, MARGIN + 30, 430, left, inner, 19)
x2 = MARGIN + cw + 30
card(d, x2, 350, cw, ch, GREEN)
T(d, x2 + 30, 380, 'Готовые модули (ЛР № 6)', F(22, True), GREEN, inner)
bullets(d, x2 + 30, 430, right, inner, 19)
ey = 350 + ch + 30
card(d, MARGIN, ey, W - 2 * MARGIN, 110, PURP)
T(d, MARGIN + 30, ey + 28, 'Принципы', F(22, True), PURP)
T(d, MARGIN + 30, ey + 66, 'Единая палитра из ЛР № 3 · статусные бейджи · семестровая навигация 1–8 на каждой странице плана · деградация без CDN', F(19), MUT, W - 2 * MARGIN - 60)
im.save(f'{OUT}/slide_10_interactive_ux.png')

# ----------------------------------------------------------------------------- S11 security
im, d = new_slide(11); head(d, 'Тестирование и безопасность', 'Защита данных и протоколирование действий')
left = ['Prepared Statements — исключены SQL-инъекции', 'htmlspecialchars / html.escape — защита от XSS',
        'FOREIGN KEYS — ссылочная целостность (0 нарушений)', 'PRAGMA foreign_key_check — контроль при каждом старте']
right = ['7 записей: время, пользователь, роль, действие, статус, IP', 'Живой поиск по журналу (client-side)',
         'Фиксация валидаций, виз ЭЦП и служебных записок', 'Статусные бейджи: УСПЕШНО / ОШИБКА / ОБНОВЛЕНО']
cw = (W - 2 * MARGIN - 30) // 2; inner = cw - 60
ch = 100 + max(bullets_height(left, inner, 20, 58), bullets_height(right, inner, 20, 58)) + 16
card(d, MARGIN, 180, cw, ch, GREEN)
T(d, MARGIN + 30, 212, 'Безопасность', F(24, True), GREEN, inner)
bullets(d, MARGIN + 30, 270, left, inner, 20, step=58)
x2 = MARGIN + cw + 30
card(d, x2, 180, cw, ch, ACC)
T(d, x2 + 30, 212, 'Журнал аудита', F(24, True), SKY, inner)
bullets(d, x2 + 30, 270, right, inner, 20, step=58)
ty = 180 + ch + 40
tests = [('15 / 15', 'страниц отвечают HTTP 200', GREEN), ('3 / 3', 'REST-эндпоинта возвращают JSON', SKY),
         ('5 / 5', 'критериев ФГОС подтверждены', AMBER), ('4 / 4', 'подписей SHA-256 верифицированы', PURP)]
n = len(tests); gap = 25; cw4 = (W - 2 * MARGIN - gap * (n - 1)) // n; x = MARGIN
for v, lab, c in tests:
    card(d, x, ty, cw4, 190, c)
    T(d, x + 30, ty + 40, v, F(48, True), c, cw4 - 60)
    para(d, x + 30, ty + 120, lab, F(18), MUT, cw4 - 60, lh=24, maxy=ty + 190)
    x += cw4 + gap
para(d, MARGIN, ty + 190 + 36, 'Все 15 страниц и 3 REST-эндпоинта отвечают HTTP 200 · оффлайн-деградация CDN-модулей без потери данных (проверка typeof перед вызовом библиотек)', F(19), MUT, W - 2 * MARGIN)
im.save(f'{OUT}/slide_11_testing_security.png')

# ----------------------------------------------------------------------------- S12 conclusion
im, d = new_slide(12); head(d, 'Заключение и результаты', 'Портал Центра АНОК готов к демонстрации и защите')
items = ['Полный цикл ЛР № 1–7: требования → проектирование → дизайн → разделы → инфоблоки → модули → защита',
         'Портал: 15 разделов + REST API, 13 таблиц БД с FOREIGN KEYS, 51 дисциплина (240 з.е., баланс 30/1080)',
         'Экспертиза ФГОС 3++: 5/5 критериев соответствуют; маршрут ЭЦП 4 этапа с верификацией SHA-256',
         '9 готовых модулей: Chart.js, FullCalendar, DataTables, SheetJS, jsPDF, CryptoJS, SweetAlert2, Toastify, Diff2Html',
         'Комплект документов: 7 отчетов (MD/DOCX/PDF), пояснительная записка, презентация 12 слайдов']
y = bullets(d, 80, 190, items, W - 160, 24, step=76)
card(d, 80, y + 10, W - 160, 190, ACC)
T(d, 110, y + 54, 'Итог: требования варианта № 15 выполнены полностью', F(28, True), SKY, W - 220)
T(d, 110, y + 118, 'Москва, 2026 · студент группы П-41 Соловьёв А.С.', F(20), MUT, W - 220)
im.save(f'{OUT}/slide_12_conclusion_impact.png')
print('slides done')

# ----------------------------------------------------------------------------- PPTX + PDF
SLIDES = sorted(f for f in os.listdir(OUT) if f.startswith('slide_') and f.endswith('.png'))
assert len(SLIDES) == 12, SLIDES
REP = f'{ROOT}/reports/lab7'

try:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9, matches 1920x1080 images
    blank = prs.slide_layouts[6]
    for f in SLIDES:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(f'{OUT}/{f}', 0, 0, prs.slide_width, prs.slide_height)
    prs.core_properties.title = 'Корпоративный портал университета — Модуль Центра АНОК (Вариант 15)'
    prs.core_properties.author = 'Соловьёв А.С., группа П-41'
    prs.save(f'{REP}/LR7_Presentation_Variant15.pptx')
    print('pptx done')
except ImportError:
    print('python-pptx not installed — PPTX not rebuilt')

imgs = [Image.open(f'{OUT}/{f}').convert('RGB') for f in SLIDES]
imgs[0].save(f'{REP}/LR7_Presentation_Variant15.pdf', save_all=True, append_images=imgs[1:], resolution=72.0,
             title='Корпоративный портал университета — Модуль Центра АНОК (Вариант 15)', author='Соловьёв А.С., группа П-41')
print('pdf done')
