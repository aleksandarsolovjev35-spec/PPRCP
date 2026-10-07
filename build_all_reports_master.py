import os
import re
from xml.sax.saxutils import escape as xml_escape

import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from PIL import Image as PILImage

import reportlab
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table as RLTable, TableStyle, PageBreak, CondPageBreak, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Register Cyrillic Fonts for ReportLab
pdfmetrics.registerFont(TTFont('DejaVuSerif', '/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'))
pdfmetrics.registerFont(TTFont('DejaVuSerif-Bold', '/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf'))
pdfmetrics.registerFont(TTFont('DejaVuSerif-Italic', '/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'))
pdfmetrics.registerFont(TTFont('DejaVuSans', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('DejaVuSansMono', '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'))

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            if self._pageNumber > 1:
                self.setFont('DejaVuSerif', 10)
                text = f'{self._pageNumber}'
                self.drawCentredString(A4[0] / 2.0, 15 * 2.83465, text)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

def add_docx_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def latex_to_text(expression):
    """Convert the small LaTeX subset used in the reports to readable Unicode.

    Markdown viewers render these expressions as mathematics.  DOCX and
    ReportLab do not, so leaving the source intact exposes strings such as
    ``$\\le 5$`` and ``\\text{credits}`` to the reader.
    """
    replacements = {
        r'\rightarrow': '→',
        r'\leftarrow': '←',
        r'\bowtie': '⋈',
        r'\leq': '≤',
        r'\le': '≤',
        r'\geq': '≥',
        r'\ge': '≥',
        r'\sum': 'Σ',
        r'\times': '×',
        r'\pm': '±',
    }
    text = expression
    # \text{...} is plain text in an exported office document.
    text = re.sub(r'\\text\{([^{}]*)\}', r'\1', text)
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = text.replace(r'\,', ' ').replace(r'\ ', ' ')
    # Do not leak unsupported command syntax into generated documents.
    text = re.sub(r'\\([A-Za-z]+)', r'\1', text)
    text = text.replace('{', '').replace('}', '')
    return re.sub(r'\s+', ' ', text).strip()


def clean_inline_md(text):
    """Normalize inline HTML/math while retaining Markdown style markers."""
    text = text.replace('&nbsp;', ' ')
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    # Also tolerate a typo present in an early version of the source reports.
    text = text.replace('$ightarrow$', '→')
    return re.sub(r'\$([^$\n]+)\$', lambda m: latex_to_text(m.group(1)), text)


def inline_md_parts(text):
    """Yield ``(text, styles)`` tuples for bold, italic and inline code.

    A small stateful parser is used instead of one regular expression because
    report sources contain nested markup such as ``**`table_name`:**``.
    """
    i = 0
    bold = False
    italic = False
    while i < len(text):
        if text.startswith('**', i) and (bold or text.find('**', i + 2) != -1):
            bold = not bold
            i += 2
            continue
        if text[i] == '*' and (italic or text.find('*', i + 1) != -1):
            italic = not italic
            i += 1
            continue
        if text[i] == '`':
            closing = text.find('`', i + 1)
            if closing != -1:
                styles = {'code'}
                if bold:
                    styles.add('bold')
                if italic:
                    styles.add('italic')
                yield text[i + 1:closing], frozenset(styles)
                i = closing + 1
                continue

        next_positions = [position for position in (
            text.find('**', i), text.find('*', i), text.find('`', i)
        ) if position != -1]
        closing = min(next_positions) if next_positions else len(text)
        if closing == i:  # unmatched marker: preserve it as ordinary text
            closing += 1
        styles = set()
        if bold:
            styles.add('bold')
        if italic:
            styles.add('italic')
        yield text[i:closing], frozenset(styles)
        i = closing


def plain_inline_md(text):
    return ''.join(value for value, _styles in inline_md_parts(clean_inline_md(text)))


def add_docx_inline(paragraph, text, font_size, default_bold=False):
    """Add inline Markdown to a Word paragraph without exposing its markers."""
    for value, styles in inline_md_parts(clean_inline_md(text)):
        run = paragraph.add_run(value)
        run.font.name = 'Courier New' if 'code' in styles else 'Times New Roman'
        run.font.size = Pt(font_size)
        run.bold = default_bold or 'bold' in styles
        run.italic = 'italic' in styles


def reportlab_inline(text):
    """Render safe ReportLab paragraph markup from inline Markdown."""
    rendered = []
    for value, styles in inline_md_parts(clean_inline_md(text)):
        value = xml_escape(value).replace('\n', '<br/>')
        # DejaVu Serif has no U+22C8; use the registered Sans face for joins.
        value = value.replace('⋈', '<font face="DejaVuSans">⋈</font>')
        if 'code' in styles:
            value = f'<font face="DejaVuSansMono">{value}</font>'
        if 'italic' in styles:
            value = f'<i>{value}</i>'
        if 'bold' in styles:
            value = f'<b>{value}</b>'
        rendered.append(value)
    return ''.join(rendered)


def is_report_section_heading(level, heading):
    """Return whether a Markdown heading is a report-level section.

    Reports use ``##`` for numbered sections and ``###`` for subsections.  A
    top-level number is followed by whitespace (``2. Section``), while a
    nested number is not (``2.1. Subsection``).  The explanatory note also
    uses an unnumbered introduction heading.
    """
    if level != 2:
        return False
    text = plain_inline_md(heading).strip()
    return bool(
        re.match(r'^\d+\.(?!\d)\s+', text)
        or re.match(r'^(?:ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|АННОТАЦИЯ|РЕФЕРАТ)\b', text, re.IGNORECASE)
    )


def is_first_content_heading(level, heading):
    """Find the first real section after the report's title-page metadata."""
    if level > 2:
        return False
    text = plain_inline_md(heading).strip()
    return bool(
        re.match(r'^\d+\.(?!\d)\s+', text)
        or re.match(r'^(?:ВВЕДЕНИЕ|ОПИСАНИЕ|ЦЕЛЬ|АННОТАЦИЯ|РЕФЕРАТ)\b', text, re.IGNORECASE)
    )


def parse_markdown_blocks(md_text, base_dir):
    lines = md_text.splitlines()
    blocks = []
    i = 0
    n = len(lines)
    
    in_code = False
    code_lines = []
    
    while i < n:
        line = lines[i]
        
        # Code block
        if line.strip().startswith('```'):
            if in_code:
                in_code = False
                blocks.append(('code', '\n'.join(code_lines)))
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
            
        if in_code:
            code_lines.append(line)
            i += 1
            continue
            
        # Empty line
        if not line.strip():
            i += 1
            continue
            
        # Divider
        if line.strip() in ['---', '***', '___']:
            i += 1
            continue
            
        # Heading
        m_h = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m_h:
            level = len(m_h.group(1))
            htext = m_h.group(2).strip()
            blocks.append(('heading', (level, htext)))
            i += 1
            continue
            
        # Image
        m_img = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if m_img:
            caption = m_img.group(1).strip()
            rel_path = m_img.group(2).strip()
            full_path = os.path.normpath(os.path.join(base_dir, rel_path))
            blocks.append(('image', (caption, full_path)))
            i += 1
            continue
            
        # Table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip('|').split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                rows.append(cells)
            if rows:
                blocks.append(('table', rows))
            continue
            
        # List item
        m_list = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.*)$', line)
        if m_list:
            indent = len(m_list.group(1))
            bullet = m_list.group(2)
            item_text = m_list.group(3).strip()
            i += 1
            while i < n and lines[i].strip() and not re.match(r'^(\s*)([-*+]|\d+\.)\s+', lines[i]) and not lines[i].startswith('#') and not lines[i].startswith('|') and not lines[i].startswith('!') and not lines[i].startswith('```'):
                item_text += ' ' + lines[i].strip()
                i += 1
            blocks.append(('list_item', (bullet, item_text, indent)))
            continue
            
        # Regular paragraph
        p_lines = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not lines[i].startswith('#') and not lines[i].startswith('|') and not lines[i].startswith('!') and not lines[i].startswith('```') and not lines[i].strip() in ['---', '***'] and not re.match(r'^(\s*)([-*+]|\d+\.)\s+', lines[i]):
            p_lines.append(lines[i].strip())
            i += 1
        blocks.append(('paragraph', ' '.join(p_lines)))
        
    return blocks

def generate_docx(doc_type, title, discipline, variant, student, city_year, blocks, output_path):
    doc = docx.Document()
    
    # Page setup (A4: 210 x 297 mm, Margins: Top 20, Bottom 20, Left 30, Right 15 mm)
    section = doc.sections[0]
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    
    # Configure default Normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    
    # Title Page
    for _ in range(8):
        doc.add_paragraph()
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r1 = p_title.add_run(doc_type.upper() + "\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)
    r1.bold = True
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p_sub.add_run(f"по дисциплине: «{discipline}»\n\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    
    r3 = p_sub.add_run(f"Тема: «{title}»\n{variant}")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(14)
    r3.bold = True
    
    for _ in range(5):
        doc.add_paragraph()
        
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.line_spacing = 1.15
    r4 = p_meta.add_run(f"Выполнил:\n{student}\n")
    r4.font.name = 'Times New Roman'
    r4.font.size = Pt(13)
    
    for _ in range(4):
        doc.add_paragraph()
        
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r5 = p_bot.add_run(city_year)
    r5.font.name = 'Times New Roman'
    r5.font.size = Pt(12)
    
    # Page break for main content
    doc.add_page_break()
    
    # Different first page header/footer
    section.different_first_page_header_footer = True
    footer = section.footer
    p_ft = footer.paragraphs[0]
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_num = p_ft.add_run()
    run_num.font.name = 'Times New Roman'
    run_num.font.size = Pt(11)
    add_docx_page_number(run_num)
    
    fig_count = 0
    tbl_count = 0
    skip_leading = True
    main_section_seen = False
    
    for b_type, b_content in blocks:
        if skip_leading:
            if b_type == 'heading' and is_first_content_heading(b_content[0], b_content[1]):
                skip_leading = False
            elif b_type == 'heading' and b_content[0] <= 2:
                continue
            elif b_type == 'paragraph' and ('Дисциплина:' in b_content or 'Вариант' in b_content or 'Тема:' in b_content):
                continue
            else:
                skip_leading = False

        if b_type == 'heading':
            level, htext = b_content
            is_main_section = is_report_section_heading(level, htext)
            p_h = doc.add_paragraph()
            p_h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_h.paragraph_format.keep_with_next = True
            p_h.paragraph_format.keep_together = True
            p_h.paragraph_format.widow_control = True
            if is_main_section and main_section_seen:
                p_h.paragraph_format.page_break_before = True
            if is_main_section:
                main_section_seen = True
            clean_h = plain_inline_md(htext)
            if level == 1:
                p_h.paragraph_format.space_before = Pt(14)
                p_h.paragraph_format.space_after = Pt(6)
                r = p_h.add_run(clean_h.upper())
                r.font.name = 'Times New Roman'
                r.font.size = Pt(16)
                r.bold = True
            elif level == 2:
                p_h.paragraph_format.space_before = Pt(12)
                p_h.paragraph_format.space_after = Pt(6)
                r = p_h.add_run(clean_h)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(14)
                r.bold = True
            else:
                p_h.paragraph_format.space_before = Pt(8)
                p_h.paragraph_format.space_after = Pt(4)
                r = p_h.add_run(clean_h)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
                r.bold = True
                
        elif b_type == 'paragraph':
            text = clean_inline_md(b_content)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Mm(12.5)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.widow_control = True
            
            add_docx_inline(p, text, 14)
                    
        elif b_type == 'list_item':
            bullet, item_text, indent = b_content
            text = clean_inline_md(item_text)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Mm(12.5 + indent * 2.5)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_together = True
            p.paragraph_format.widow_control = True
            
            prefix = "• " if bullet in ['-', '*', '+'] else f"{bullet} "
            r_pre = p.add_run(prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(14)
            r_pre.bold = (bullet not in ['-', '*', '+'])
            
            add_docx_inline(p, text, 14)
                    
        elif b_type == 'image':
            caption, img_path = b_content
            if os.path.exists(img_path):
                fig_count += 1
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.paragraph_format.keep_with_next = True
                p_img.paragraph_format.keep_together = True
                
                with PILImage.open(img_path) as im:
                    w, h = im.size
                
                max_w_mm = 160
                max_h_mm = 220
                scale = min(max_w_mm / w, max_h_mm / h)
                picture_width = Mm(w * scale)
                picture_height = Mm(h * scale)
                p_img.add_run().add_picture(img_path, width=picture_width, height=picture_height)
                
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(4)
                p_cap.paragraph_format.space_after = Pt(12)
                p_cap.paragraph_format.keep_together = True
                p_cap.paragraph_format.widow_control = True
                cap_text = f"Рисунок {fig_count} — {plain_inline_md(caption)}" if caption else f"Рисунок {fig_count}"
                r_cap = p_cap.add_run(cap_text)
                r_cap.font.name = 'Times New Roman'
                r_cap.font.size = Pt(12)
                r_cap.italic = True
                
        elif b_type == 'table':
            rows = b_content
            if rows:
                tbl_count += 1
                p_tcap = doc.add_paragraph()
                p_tcap.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_tcap.paragraph_format.first_line_indent = Mm(12.5)
                p_tcap.paragraph_format.space_before = Pt(8)
                p_tcap.paragraph_format.space_after = Pt(4)
                p_tcap.paragraph_format.keep_with_next = True
                p_tcap.paragraph_format.keep_together = True
                r_tcap = p_tcap.add_run(f"Таблица {tbl_count} — Спецификация данных")
                r_tcap.font.name = 'Times New Roman'
                r_tcap.font.size = Pt(12)
                r_tcap.bold = True
                
                num_rows = len(rows)
                num_cols = max(len(r) for r in rows)
                
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table)
                
                for r_idx, row_cells in enumerate(rows):
                    row = table.rows[r_idx]
                    tr_pr = row._tr.get_or_add_trPr()
                    cant_split = OxmlElement('w:cantSplit')
                    tr_pr.append(cant_split)
                    if r_idx == 0:
                        repeat_header = OxmlElement('w:tblHeader')
                        repeat_header.set(qn('w:val'), 'true')
                        tr_pr.append(repeat_header)
                    for c_idx, cell_value in enumerate(row_cells):
                        if c_idx < num_cols:
                            cell = row.cells[c_idx]
                            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                            if r_idx == 0:
                                set_cell_shading(cell, "E2E8F0") # Slate-200
                            cell_text = clean_inline_md(cell_value)
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.05
                            
                            add_docx_inline(p, cell_text, 11, default_bold=(r_idx == 0))
                                        
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                
        elif b_type == 'code':
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Mm(10)
            p_code.paragraph_format.space_before = Pt(6)
            p_code.paragraph_format.space_after = Pt(6)
            r = p_code.add_run(b_content)
            r.font.name = 'Courier New'
            r.font.size = Pt(10)
            
    doc.save(output_path)
    print(f"  [DOCX] Successfully saved: {output_path}")

def generate_pdf(doc_type, title, discipline, variant, student, city_year, blocks, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=30 * 2.83465,
        rightMargin=15 * 2.83465,
        topMargin=20 * 2.83465,
        bottomMargin=20 * 2.83465
    )
    
    styles = getSampleStyleSheet()
    
    title_top_style = ParagraphStyle(
        'TitleTop',
        fontName='DejaVuSerif-Bold',
        fontSize=10,
        leading=14,
        alignment=1,
        spaceAfter=15
    )
    
    title_main_style = ParagraphStyle(
        'TitleMain',
        fontName='DejaVuSerif-Bold',
        fontSize=15,
        leading=20,
        alignment=1,
        spaceAfter=10
    )
    
    title_sub_style = ParagraphStyle(
        'TitleSub',
        fontName='DejaVuSerif',
        fontSize=12,
        leading=16,
        alignment=1,
        spaceAfter=15
    )
    
    title_meta_style = ParagraphStyle(
        'TitleMeta',
        fontName='DejaVuSerif',
        fontSize=11,
        leading=15,
        alignment=2,
        spaceAfter=15
    )
    
    title_bot_style = ParagraphStyle(
        'TitleBot',
        fontName='DejaVuSerif',
        fontSize=11,
        leading=14,
        alignment=1,
        spaceAfter=0
    )
    
    h1_style = ParagraphStyle(
        'H1',
        fontName='DejaVuSerif-Bold',
        fontSize=14,
        leading=18,
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'H2',
        fontName='DejaVuSerif-Bold',
        fontSize=12,
        leading=16,
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )
    
    h3_style = ParagraphStyle(
        'H3',
        fontName='DejaVuSerif-Bold',
        fontSize=11,
        leading=15,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'Body',
        fontName='DejaVuSerif',
        fontSize=11,
        leading=15,
        alignment=4,
        firstLineIndent=1.25 * 28.3465,
        spaceAfter=6,
        allowWidows=0,
        allowOrphans=0
    )
    
    list_style = ParagraphStyle(
        'List',
        fontName='DejaVuSerif',
        fontSize=11,
        leading=15,
        alignment=4,
        leftIndent=1.25 * 28.3465,
        spaceAfter=3,
        allowWidows=0,
        allowOrphans=0
    )
    
    caption_style = ParagraphStyle(
        'Caption',
        fontName='DejaVuSerif',
        fontSize=10,
        leading=13,
        alignment=1,
        spaceBefore=4,
        spaceAfter=10
    )
    
    table_cell_style = ParagraphStyle(
        'TableCell',
        fontName='DejaVuSerif',
        fontSize=9,
        leading=12,
        alignment=0
    )
    
    table_head_style = ParagraphStyle(
        'TableHead',
        fontName='DejaVuSerif-Bold',
        fontSize=9.5,
        leading=13,
        alignment=1
    )
    
    # A report-level section starts on a fresh page, but do not insert a
    # blank page when the previous flowable already ended at the top of one.
    doc._calc()
    frame_height = doc.height - 12  # ReportLab's default top/bottom frame padding.
    section_break_height = max(1, frame_height - 12)

    story = []
    
    # Title Page
    story.append(Spacer(1, 160))
    story.append(Paragraph(xml_escape(doc_type.upper()), title_main_style))
    story.append(Paragraph(f"по дисциплине: «{xml_escape(discipline)}»<br/><br/><b>Тема: «{xml_escape(title)}»</b><br/><b>{xml_escape(variant)}</b>", title_sub_style))
    story.append(Spacer(1, 60))
    
    student_fmt = xml_escape(student).replace('\n', '<br/>')
    story.append(Paragraph(f"<b>Выполнил:</b><br/>{student_fmt}", title_meta_style))
    story.append(Spacer(1, 60))
    story.append(Paragraph(xml_escape(city_year), title_bot_style))
    story.append(PageBreak())
    
    # Main Content
    fig_count = 0
    tbl_count = 0
    skip_leading = True
    main_section_seen = False
    
    for b_type, b_content in blocks:
        if skip_leading:
            if b_type == 'heading' and is_first_content_heading(b_content[0], b_content[1]):
                skip_leading = False
            elif b_type == 'heading' and b_content[0] <= 2:
                continue
            elif b_type == 'paragraph' and ('Дисциплина:' in b_content or 'Вариант' in b_content or 'Тема:' in b_content):
                continue
            else:
                skip_leading = False
                
        if b_type == 'heading':
            level, htext = b_content
            is_main_section = is_report_section_heading(level, htext)
            if is_main_section and main_section_seen:
                story.append(CondPageBreak(section_break_height))
            if is_main_section:
                main_section_seen = True
            htext_fmt = reportlab_inline(htext)
            if level == 1:
                story.append(Paragraph(htext_fmt.upper(), h1_style))
            elif level == 2:
                story.append(Paragraph(htext_fmt, h2_style))
            else:
                story.append(Paragraph(htext_fmt, h3_style))
                
        elif b_type == 'paragraph':
            story.append(Paragraph(reportlab_inline(b_content), body_style))
            
        elif b_type == 'list_item':
            bullet, item_text, indent = b_content
            prefix = "• " if bullet in ['-', '*', '+'] else f"{bullet} "
            story.append(Paragraph(f"<b>{xml_escape(prefix)}</b>{reportlab_inline(item_text)}", list_style))
            
        elif b_type == 'image':
            caption, img_path = b_content
            if os.path.exists(img_path):
                fig_count += 1
                with PILImage.open(img_path) as im:
                    w, h = im.size
                
                max_w_pt = 165 * 2.83465
                max_h_pt = 220 * 2.83465
                
                scale = min(max_w_pt / w, max_h_pt / h, 1.0)
                img_w = w * scale
                img_h = h * scale
                
                cap_text = f"<i>Рисунок {fig_count} — {reportlab_inline(caption)}</i>" if caption else f"<i>Рисунок {fig_count}</i>"
                img_flowable = RLImage(img_path, width=img_w, height=img_h)
                img_flowable.keepWithNext = True
                cap_flowable = Paragraph(cap_text, caption_style)
                story.extend([img_flowable, cap_flowable])
                
        elif b_type == 'table':
            rows = b_content
            if rows:
                tbl_count += 1
                tcap_style = ParagraphStyle(
                    f'TCap{tbl_count}',
                    fontName='DejaVuSerif-Bold',
                    fontSize=10,
                    spaceBefore=8,
                    spaceAfter=4,
                    firstLineIndent=1.25 * 28.3465,
                    keepWithNext=True
                )
                tcap = Paragraph(f"<b>Таблица {tbl_count} — Спецификация данных</b>", tcap_style)
                
                table_data = []
                for r_idx, r in enumerate(rows):
                    row_data = []
                    for c in r:
                        c_fmt = reportlab_inline(c)
                        if r_idx == 0:
                            row_data.append(Paragraph(c_fmt, table_head_style))
                        else:
                            row_data.append(Paragraph(c_fmt, table_cell_style))
                    table_data.append(row_data)
                    
                num_cols = max(len(r) for r in rows)
                total_w = 165 * 2.83465
                col_w = total_w / num_cols
                
                t = RLTable(table_data, colWidths=[col_w]*num_cols, repeatRows=1, splitByRow=1)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                    ('TOPPADDING', (0,0), (-1,-1), 4),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                    ('LEFTPADDING', (0,0), (-1,-1), 4),
                    ('RIGHTPADDING', (0,0), (-1,-1), 4),
                ]))
                # Keep the table caption, heading immediately before it, and
                # the table together as a flow chain, while still allowing a
                # long table to continue onto following pages.
                story.extend([tcap, t, Spacer(1, 8)])
                
        elif b_type == 'code':
            story.append(Preformatted(b_content, ParagraphStyle('Code', fontName='DejaVuSansMono', fontSize=8.5, leading=11, spaceBefore=4, spaceAfter=6, leftIndent=10)))
            
    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"  [PDF]  Successfully saved: {output_path}")

DOCUMENTS = [
    {
        "dir": "reports/lab1",
        "md": "reports/lab1/LR1_Report_Variant15.md",
        "docx": "reports/lab1/LR1_Report_Variant15.docx",
        "pdf": "reports/lab1/LR1_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 1",
        "title": "Выявление требований к порталу и составление их формализованного описания. Выбор инструментального средства для создания портала",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab2",
        "md": "reports/lab2/LR2_Report_Variant15.md",
        "docx": "reports/lab2/LR2_Report_Variant15.docx",
        "pdf": "reports/lab2/LR2_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 2",
        "title": "Анализ и проектирование портала. Разворачивание программного обеспечения",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab3",
        "md": "reports/lab3/LR3_Report_Variant15.md",
        "docx": "reports/lab3/LR3_Report_Variant15.docx",
        "pdf": "reports/lab3/LR3_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 3",
        "title": "Создание шаблона дизайна портала и главных пользовательских интерфейсов",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab4",
        "md": "reports/lab4/LR4_Report_Variant15.md",
        "docx": "reports/lab4/LR4_Report_Variant15.docx",
        "pdf": "reports/lab4/LR4_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 4",
        "title": "Создание структуры разделов и страниц портала",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab5",
        "md": "reports/lab5/LR5_Report_Variant15.md",
        "docx": "reports/lab5/LR5_Report_Variant15.docx",
        "pdf": "reports/lab5/LR5_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 5",
        "title": "Работа с инфоблоками и информационное наполнение портала",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab6",
        "md": "reports/lab6/LR6_Report_Variant15.md",
        "docx": "reports/lab6/LR6_Report_Variant15.docx",
        "pdf": "reports/lab6/LR6_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 6",
        "title": "Работа с готовыми модулями и компонентами",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab7",
        "md": "reports/lab7/LR7_Report_Variant15.md",
        "docx": "reports/lab7/LR7_Report_Variant15.docx",
        "pdf": "reports/lab7/LR7_Report_Variant15.pdf",
        "doc_type": "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 7",
        "title": "Подготовка презентации и пояснительной записки",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    },
    {
        "dir": "reports/lab7",
        "md": "reports/lab7/LR7_Explanatory_Note_Variant15.md",
        "docx": "reports/lab7/LR7_Explanatory_Note_Variant15.docx",
        "pdf": "reports/lab7/LR7_Explanatory_Note_Variant15.pdf",
        "doc_type": "ПОЯСНИТЕЛЬНАЯ ЗАПИСКА К ВЫПУСКНОМУ ПРОЕКТУ",
        "title": "Проектирование и разработка корпоративного веб-портала университета — Модуль Центра аккредитации и независимой оценки качества образования (АНОК)",
        "discipline": "Практикум по разработке корпоративного портала",
        "variant": "Вариант № 15: «Университет. Центр аккредитации и независимой оценки качества образования (АНОК)»",
        "student": "студент группы П-41\nСоловьёв А.С.",
        "city_year": "Москва, 2026 г."
    }
]

def main():
    print("Starting regeneration of all 8 standardized documents (without supervisor block)...")
    for idx, doc_meta in enumerate(DOCUMENTS):
        print(f"\n[{idx+1}/8] Processing: {doc_meta['md']}")
        with open(doc_meta['md'], 'r', encoding='utf-8') as f:
            md_text = f.read()
        blocks = parse_markdown_blocks(md_text, doc_meta['dir'])
        print(f"  Parsed {len(blocks)} content blocks.")
        generate_docx(
            doc_type=doc_meta['doc_type'],
            title=doc_meta['title'],
            discipline=doc_meta['discipline'],
            variant=doc_meta['variant'],
            student=doc_meta['student'],
            city_year=doc_meta['city_year'],
            blocks=blocks,
            output_path=doc_meta['docx']
        )
        generate_pdf(
            doc_type=doc_meta['doc_type'],
            title=doc_meta['title'],
            discipline=doc_meta['discipline'],
            variant=doc_meta['variant'],
            student=doc_meta['student'],
            city_year=doc_meta['city_year'],
            blocks=blocks,
            output_path=doc_meta['pdf']
        )

    print("\nAll 8 standardized DOCX and PDF documents regenerated successfully without supervisor block!")


if __name__ == '__main__':
    main()
